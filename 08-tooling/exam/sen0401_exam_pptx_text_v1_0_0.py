"""Shared reader for the critical-points work: slide text, notes and legacy diagram text.
Read only; never writes a deck. Version 1.0.0."""
import glob, re, sys
from pptx import Presentation

ROOT = '/home/claude/sen0401/'
MAT = ROOT + '03-materials/'
LEG = MAT + 'owner-legacy-2025/slides-2025/'


def norm(t):
    t = t.replace('’', "'").replace('‘', "'").replace('“', '"').replace('”', '"')
    t = t.replace(' ', ' ')
    return re.sub(r'\s+', ' ', t).strip()


def _ver(p):
    m = re.search(r'_v(\d+)_(\d+)_(\d+)\.pptx$', p)
    return tuple(int(x) for x in m.groups())


def newest_deck(ch):
    fs = sorted(glob.glob(MAT + 'ch%02d/*.pptx' % ch), key=_ver)
    return fs[-1]


def slides_of(path):
    """list of dicts {n, slide, notes} with normalised text"""
    pr = Presentation(path)
    out = []
    for i, s in enumerate(pr.slides, 1):
        parts = []
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                parts.append(sh.text_frame.text)
            if getattr(sh, 'has_table', False) and sh.has_table:
                for r in sh.table.rows:
                    for c in r.cells:
                        parts.append(c.text)
        notes = ''
        if s.has_notes_slide:
            notes = s.notes_slide.notes_text_frame.text
        out.append({'n': i, 'slide': norm(' '.join(parts)), 'notes': norm(notes)})
    return out


def legacy_diagram_text(path, n):
    """text of the SmartArt parts of a legacy slide (diagram data) plus its title"""
    pr = Presentation(path)
    s = pr.slides[n - 1]
    txt = []
    for sh in s.shapes:
        if sh.has_text_frame:
            txt.append(sh.text_frame.text)
    for r in s.part.rels.values():
        if r.reltype.endswith('diagramData'):
            txt.append(' '.join(re.findall(r'<a:t>([^<]*)</a:t>', r.target_part.blob.decode('utf8', 'ignore'))))
    return norm(' '.join(txt))
