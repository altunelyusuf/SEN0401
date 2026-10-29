#!/usr/bin/env python3
"""sen0401_deck_teach_measure v1_0_0 - measured facts on the code and the teaching text of the SEN0401 decks of chapters 1 to 6 (owner request of 2026-10-30 02:21 Istanbul on the chapter 6 deck). Facts only; nothing is edited. Per deck: slides; code lines (prompt rows with the prompt stripped); code lines with leading indentation; one-line compound statements (a for/if/while/def/with/class/try header with its body on the same line); lines with the word lambda; text shapes that begin with a 5N1K label word (WHAT: WHO: WHEN: WHERE: WHY: HOW:); speaker-note words (average, maximum). Usage: measure.py [--json FILE] | --selftest"""
import glob, json, re, sys
sys.dont_write_bytecode = True
R = "/home/claude/sen0401/03-materials/"
PROMPT = re.compile(r"^(>>> |\.\.\. |\.\.\.$)")
LABEL = re.compile(r"^\s*(WHAT|WHO|WHEN|WHERE|WHY|HOW)\s*[:\-–—]", re.I)
COMPOUND = re.compile(r"^\s*(for|while|if|elif|else|def|class|with|try|except|finally)\b[^#]*:\s*\S")
def latest(pat):
    return sorted(glob.glob(pat), key=lambda x: [int(v) for v in re.findall(r"_v(\d+)_(\d+)_(\d+)", x)[-1]])[-1]
def one_line_compound(l):
    if not COMPOUND.match(l): return False
    head = l.split("#")[0]
    # a colon followed by text outside brackets/strings is a body on the header line; slices and dict colons sit inside brackets
    depth = 0; q = None
    for i, c in enumerate(head):
        if q:
            if c == q: q = None
        elif c in "'\"": q = c
        elif c in "([{": depth += 1
        elif c in ")]}": depth -= 1
        elif c == ":" and depth == 0 and head[i + 1:].strip(): return True
    return False
def selftest():
    ok = one_line_compound("for i in range(3): print(i)") and not one_line_compound("for i in range(3):") and not one_line_compound("x = d['a']") and not one_line_compound("if x[1:2] == y:") and one_line_compound("if x: y = 1")
    ok &= bool(LABEL.match("WHAT: a block")) and not LABEL.match("What a block is") and bool(LABEL.match("WHO - Satoshi"))
    print("SELFTEST", "PASS" if ok else "FAIL"); return ok
def main():
    from pptx import Presentation
    out = {}
    for n in range(1, 7):
        deck = latest(R + "ch%02d/SEN0401_Ch%02d_*_3e_v*.pptx" % (n, n)); pr = Presentation(deck)
        code = []; lab = []; notes = []; lam = 0
        for i, s in enumerate(pr.slides, 1):
            for sh in s.shapes:
                if not sh.has_text_frame: continue
                for p in sh.text_frame.paragraphs:
                    for ln in "".join(r.text for r in p.runs).split("\n"):
                        if PROMPT.match(ln): code.append((i, PROMPT.sub("", ln, count=1)))
                        if "lambda" in ln: lam += 1
                t = sh.text_frame.text.strip()
                if LABEL.match(t): lab.append(i)
            if s.has_notes_slide: notes.append(len(s.notes_slide.notes_text_frame.text.split()))
            else: notes.append(0)
        out[n] = {"deck": deck.split("/")[-1], "slides": len(pr.slides), "code_lines_prompt_rows_stripped": len(code), "with_leading_indentation": sum(1 for _, l in code if l[:1] in " \t"),
                  "one_line_compound": sum(1 for _, l in code if one_line_compound(l)), "lambda_lines": lam, "text_shapes_beginning_with_5N1K_label": len(lab), "slides_with_such_shape": sorted(set(lab)),
                  "notes_words_average": round(sum(notes) / max(1, len(notes)), 1), "notes_words_max": max(notes or [0]), "slides_without_notes": sum(1 for x in notes if x == 0)}
    print(json.dumps(out, indent=1))
    if "--json" in sys.argv: json.dump(out, open(sys.argv[sys.argv.index("--json") + 1], "x"), indent=1)
if __name__ == "__main__": sys.exit(0 if selftest() else 1) if "--selftest" in sys.argv else main()
