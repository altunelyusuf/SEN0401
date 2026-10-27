#!/usr/bin/env python3
"""deck_charts_proof v1.0.0 (chapter 13, task TK_Deck13_Charts_Build) - prints the proofs, in sections, every result as measured (PASS / FAIL, exit 1 on any FAIL):
  sha      SHA-256 of deck 2.2.0 and manifest 1.1.0 now, against the values taken before the build (argument 1 file: lines 'sha  path')
  diff     extracted text and notes of every slide of deck 2.2.0 against 2.3.0: only slide 24, only the passphrase sentence
  parts    package parts whose bytes differ between the two decks
  links    hyperlinks (slide, target) of both decks
  run      all examples of manifest 1.2.0 run on the stored results (one fresh namespace per example after the setup)
  code     manifest code equals the '>>> ' lines of its slides in deck 2.3.0; manifest 1.2.0 differs from 1.1.0 only in manifest, deck, deck_version
  errfirst no run result carries 'Error:' in its first 160 characters
Usage: python3 -B (sha diff parts links code) or python3.14 -B (run errfirst) deck_charts_proof_v1_0_0.py BEFORE_SHA_FILE SECTION"""
import difflib, hashlib, json, os, re, sys, tempfile, zipfile, shutil
Presentation = None  # imported where needed (python3.14 has no python-pptx)
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(os.path.dirname(HERE)); D = os.path.join(ROOT, "03-materials", "ch13")
OLD = os.path.join(D, "SEN0401_Ch13_BitcoinSecurity_3e_v2_2_0.pptx"); NEW = os.path.join(D, "SEN0401_Ch13_BitcoinSecurity_3e_v2_3_0.pptx")
MOLD = os.path.join(D, "page", "chapter_code_ch13_v1_1_0.json"); MNEW = os.path.join(D, "page", "chapter_code_ch13_v1_2_0.json")
fail = []
def say(ok, msg):
    print(("PASS " if ok else "FAIL ") + msg)
    if not ok: fail.append(msg)
def sha(p): return hashlib.sha256(open(p, "rb").read()).hexdigest()
def slide_texts(path):
    from pptx import Presentation
    out = []
    for i, s in enumerate(Presentation(path).slides, 1):
        lines = []
        for sh in s.shapes:
            if sh.has_text_frame: lines += ["[%s] %s" % (sh.name, p.text) for p in sh.text_frame.paragraphs]
            if sh.has_table: lines += ["[%s] %s" % (sh.name, " | ".join(c.text for c in r.cells)) for r in sh.table.rows]
        out.append((i, lines, s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""))
    return out
def sec_sha(before):
    for l in open(before).read().strip().splitlines():
        h, p = l.split(None, 1); now = sha(os.path.join(ROOT, p.strip()))
        print("before", h, p); print("after ", now, p); say(h == now, "byte-identical " + p)
    print("new files:"); [print(" ", sha(p), os.path.relpath(p, ROOT)) for p in (NEW, MNEW)]
def sec_diff():
    a, b = slide_texts(OLD), slide_texts(NEW); say(len(a) == len(b) == 55, "slide count %d and %d" % (len(a), len(b)))
    changed = []
    for (i, ta, na), (_, tb, nb) in zip(a, b):
        if ta != tb or na != nb:
            changed.append(i)
            print("--- slide %d text/notes v2_2_0 against v2_3_0 (notes %s)" % (i, "same" if na == nb else "DIFFER"))
            for l in difflib.unified_diff(ta, tb, "v2_2_0", "v2_3_0", lineterm="", n=0): print(l)
    say(changed == [24], "slides whose text or notes differ: %s" % changed)
    gone = [l for l in a[23][1] if l not in b[23][1]]; new = [l for l in b[23][1] if l not in a[23][1]]
    say(len(gone) == 1 and len(new) == 1 and gone[0].endswith('entered wa"') and new[0].endswith('entered was incorrect."') and gone[0][:-len('entered wa"')] == new[0][:-len('entered was incorrect."')], "slide 24: exactly one line differs, only the cut message became the whole message")
    say(a[23][2] == b[23][2], "notes of slide 24 equal")
def sec_parts():
    za, zb = zipfile.ZipFile(OLD), zipfile.ZipFile(NEW); na, nb = set(za.namelist()), set(zb.namelist())
    diff = sorted(n for n in na & nb if za.read(n) != zb.read(n))
    print("parts in v2_2_0: %d, in v2_3_0: %d, parts in only one: %s" % (len(na), len(nb), sorted(na ^ nb))); print("parts whose bytes differ:", diff)
    say(na == nb and set(diff) <= {"ppt/charts/chart2.xml", "ppt/slides/slide21.xml", "ppt/slides/slide24.xml", "docProps/core.xml"} and {"ppt/charts/chart2.xml", "ppt/slides/slide24.xml", "docProps/core.xml"} <= set(diff), "differing parts are only the chart part of slide 21, the slide 21/24 parts and the properties part")
    x = zb.read("ppt/charts/chart2.xml").decode(); y = za.read("ppt/charts/chart2.xml").decode()
    v = lambda t: re.findall(r"<c:v>([^<]*)</c:v>", t)
    say(v(x) == v(y) and "<c:barChart>" in x and x.count('showVal val="1"') >= 2, "chart2: same cached values and categories (%d), still a bar chart, showVal 1 in series and group labels" % len(v(x)))
def links(path):
    z = zipfile.ZipFile(path); out = []
    for n in sorted((n for n in z.namelist() if re.match(r"ppt/slides/_rels/slide\d+\.xml\.rels$", n)), key=lambda n: int(re.findall(r"\d+", n)[0])):
        for m in re.finditer(r'Type="[^"]*/hyperlink"[^>]*Target="([^"]+)"|Target="([^"]+)"[^>]*Type="[^"]*/hyperlink"', z.read(n).decode()):
            out.append((int(re.findall(r"\d+", n)[0]), m.group(1) or m.group(2)))
    return out
def sec_links():
    a, b = links(OLD), links(NEW); print("hyperlinks v2_2_0: %d, v2_3_0: %d" % (len(a), len(b)))
    for l in b: print(" ", l)
    ua, ub = {t for _, t in a}, {t for _, t in b}
    print("distinct addresses v2_2_0: %d, v2_3_0: %d (one address, the Shamir paper, is linked on two slides)" % (len(ua), len(ub)))
    say(len(ua) == 14 and len(ub) == 14 and sorted(a) == sorted(b), "14 distinct hyperlink addresses in v2_2_0 and in v2_3_0 (%d and %d link entries), same (slide, target) entries; lost: %s" % (len(a), len(b), sorted(set(a) - set(b))))
def stored(m):
    whl = os.path.join(ROOT, m["library"]["file"]); say(sha(whl) == m["library"]["sha256"], "library wheel sha256 equals the manifest's")
    tmp = tempfile.mkdtemp(prefix="sen0401_fx_ch13_i41_"); zipfile.ZipFile(whl).extractall(tmp); return tmp
def sec_run():
    m = json.load(open(MNEW)); tmp = stored(m); sys.path.insert(0, tmp); cwd = os.getcwd(); os.chdir(tmp); ok = 0; n = 0; bad = []
    try:
        for e in m["examples"]:
            n += 1; ns = {}
            try:
                exec(m["setup"]["code"], ns); exec(e["code"], ns); good = True
                for s in e["show"]:
                    got = repr(eval(s["expression"], ns))
                    if got != s["expected"]: good = False; bad.append((e["id"], s["expression"], s["expected"], got))
            except Exception as x: good = False; bad.append((e["id"], "EXC", "", "%s: %s" % (type(x).__name__, x)))
            ok += good; print("%s %-5s slide %2d %s" % ("ok  " if good else "FAIL", e["id"], e["slide"], e["name"]))
    finally: os.chdir(cwd); shutil.rmtree(tmp, ignore_errors=True)
    print("python", sys.version.split()[0], "(expected values recorded under %s)" % m["python_of_expected_values"])
    say(ok == n == 28 and not bad, "%d of %d examples ran, every output equal to its expected output; mismatches: %s" % (ok, n, bad))
def sec_code():
    from pptx import Presentation
    m = json.load(open(MNEW)); pr = Presentation(NEW); badc = []
    for e in m["examples"]:
        lines = [p.text[4:] for sh in pr.slides[e["slide"] - 1].shapes if sh.has_text_frame for p in sh.text_frame.paragraphs if p.text.startswith(">>> ")]
        if lines != e["code"].split("\n"): badc.append((e["id"], e["slide"]))
    say(not badc and len(m["examples"]) == 28, "manifest code equals the '>>> ' code of its slide in deck 2.3.0 for %d examples; differing: %s" % (len(m["examples"]), badc))
    m0 = json.load(open(MOLD)); chg = [a["id"] for a, b in zip(m0["examples"], m["examples"]) if a != b]
    say(chg == [] and len(m0["examples"]) == len(m["examples"]), "examples that differ between manifest 1.1.0 and 1.2.0: %s" % chg)
    say(sorted(k for k in set(m0) | set(m) if m0.get(k) != m.get(k)) == ["deck", "deck_version", "manifest"] and (m["manifest"], m["deck"], m["deck_version"]) == ("chapter_code_ch13_v1_2_0", "SEN0401_Ch13_BitcoinSecurity_3e_v2_3_0.pptx", "2.3.0"), "top-level fields that differ: %s" % sorted(k for k in set(m0) | set(m) if m0.get(k) != m.get(k)))
    say([e["slide"] for e in m["examples"]] == [e["slide"] for e in m0["examples"]], "slide numbers of the manifest unchanged")
def sec_errfirst():
    m = json.load(open(MNEW)); tmp = stored(m); sys.path.insert(0, tmp); cwd = os.getcwd(); os.chdir(tmp); hits = []; n = 0
    try:
        for e in m["examples"]:
            ns = {}; exec(m["setup"]["code"], ns); exec(e["code"], ns)
            for s in e["show"]:
                n += 1
                for txt in (repr(eval(s["expression"], ns)), s["expected"]):
                    if "Error:" in txt[:160]: hits.append((e["id"], s["expression"], txt[:60]))
    finally: os.chdir(cwd); shutil.rmtree(tmp, ignore_errors=True)
    say(not hits, "%d results (run and expected text) of %d examples: 'Error:' in the first 160 characters: %s" % (n * 2, len(m["examples"]), hits))
S = {"sha": lambda: sec_sha(sys.argv[1]), "diff": sec_diff, "parts": sec_parts, "links": sec_links, "run": sec_run, "code": sec_code, "errfirst": sec_errfirst}
S[sys.argv[2]](); print("RESULT", "FAIL" if fail else "PASS"); sys.exit(1 if fail else 0)
