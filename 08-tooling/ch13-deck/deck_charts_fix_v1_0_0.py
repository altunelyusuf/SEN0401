#!/usr/bin/env python3
"""deck_charts_fix v1.0.0 (chapter 13) - task TK_Deck13_Charts_Build: makes the NEW deck 2.3.0, the NEW plan 2.3.0 and the NEW chapter-code manifest 1.2.0
from deck/plan 2.2.0 and manifest 1.1.0 (all three stay unchanged). TWO changes only:
 (1) the native chart ppt/charts/chart2.xml (slide 21, example 10) carries data labels like the five other charts (showVal 1, 9 pt, outside end, turned upright-to-vertical so 15 bars stay readable); values and chart type unchanged;
 (2) the body text of slide 24 says the whole message example 12 prints: The wallet passphrase entered was incorrect.
Manifest 1.2.0: only the fields manifest, deck, deck_version change. Usage: python3 -B deck_charts_fix_v1_0_0.py"""
__version__ = "1.0.0"
import copy, json, os, subprocess, sys
from pptx import Presentation
from lxml import etree
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(os.path.dirname(HERE)); D = os.path.join(ROOT, "03-materials", "ch13")
OLD = os.path.join(D, "SEN0401_Ch13_BitcoinSecurity_3e_v2_2_0.pptx"); NEW = os.path.join(D, "SEN0401_Ch13_BitcoinSecurity_3e_v2_3_0.pptx")
MOLD = os.path.join(D, "page", "chapter_code_ch13_v1_1_0.json"); MNEW = os.path.join(D, "page", "chapter_code_ch13_v1_2_0.json")
POLD = os.path.join(HERE, "deck_plan_v2_2_0.json"); PNEW = os.path.join(HERE, "deck_plan_v2_3_0.json")
for p in (NEW, MNEW, PNEW):
    if os.path.exists(p): sys.exit("REFUSED: %s exists" % p)
OUT = json.load(open(os.path.join(HERE, "examples_out_v2_2_0.json")))
msg = eval(OUT["ex12"][1][1]); assert msg == "The wallet passphrase entered was incorrect."
OLDBODY = 'Sending is refused with: "Please enter the wallet". The wrong passphrase gives: "The wallet passphrase entered wa"'
BODY = 'Sending is refused with: "Please enter the wallet". The wrong passphrase gives: "%s"' % msg

# ---- plan 2.3.0
plan = json.load(open(POLD))
ch = [i for i in plan["slides"][20]["items"] if i["t"] == "chart"][0]; assert ch["dataLabels"] is False; ch["dataLabels"] = True
tx = [i for i in plan["slides"][23]["items"] if i["t"] == "text" and i["body"] == OLDBODY][0]; tx["body"] = BODY
plan["meta"]["version"] = "2.3.0"; plan["meta"]["sub"] = "SEN0401, third-edition deck 2.3.0"
plan["meta"]["revision"] = "2026-10-27, task TK_Deck13_Charts_Build: the chart of slide 21 carries data labels; the body text of slide 24 says the whole message example 12 prints; nothing else changed"
json.dump(plan, open(PNEW, "w"), ensure_ascii=False, indent=1)

# ---- deck 2.3.0
pr = Presentation(OLD)
s24 = pr.slides[23]; sh = [x for x in s24.shapes if x._element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}txBody") is not None and x.text_frame.text == OLDBODY]  # never touch text_frame of a shape without a txBody (python-pptx would add one); assert len(sh) == 1
runs = sh[0].text_frame.paragraphs[0].runs; assert len(runs) == 1 and len(sh[0].text_frame.paragraphs) == 1; runs[0].text = BODY
cshape = [x for x in pr.slides[20].shapes if x.has_chart]; assert len(cshape) == 1
cs = cshape[0].chart.part._element
C = "http://schemas.openxmlformats.org/drawingml/2006/chart"; A = "http://schemas.openxmlformats.org/drawingml/2006/main"
q = lambda t: "{%s}%s" % (C, t)
vals_before = [v.text for v in cs.iter(q("v"))]
LBL = ('<c:dLbls xmlns:c="%s" xmlns:a="%s"><c:numFmt formatCode="#,##0" sourceLinked="0"/><c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>'
       '<c:txPr><a:bodyPr rot="-5400000" vert="horz"/><a:lstStyle/><a:p><a:pPr><a:defRPr b="0" i="0" strike="noStrike" sz="900" u="none"><a:solidFill><a:srgbClr val="2B2622"/></a:solidFill><a:latin typeface="Calibri"/></a:defRPr></a:pPr><a:endParaRPr lang="en-US"/></a:p></c:txPr>'
       '<c:dLblPos val="outEnd"/><c:showLegendKey val="0"/><c:showVal val="1"/><c:showCatName val="0"/><c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbls>') % (C, A)
ser = cs.findall(".//" + q("ser")); assert len(ser) == 1
old = ser[0].find(q("dLbls")); assert old is not None
ser[0].remove(old)
dpts = ser[0].findall(q("dPt")); new = etree.fromstring(LBL); dpts[-1].addnext(new)  # schema order: dPt* then dLbls
grp = cs.find(".//" + q("barChart") + "/" + q("dLbls")); assert grp is not None
for e in grp.iter(q("showVal")): e.set("val", "1")
assert [v.text for v in cs.iter(q("v"))] == vals_before
pr.save(NEW)
subprocess.check_call(["python3", os.path.join(os.path.dirname(HERE), "office_stamp_v1_0_0.py"), NEW, "2.3.0", "SEN0401 Chapter 13"])

# ---- manifest 1.2.0
m = json.load(open(MOLD)); m["manifest"] = "chapter_code_ch13_v1_2_0"; m["deck"] = "SEN0401_Ch13_BitcoinSecurity_3e_v2_3_0.pptx"; m["deck_version"] = "2.3.0"
json.dump(m, open(MNEW, "w"), ensure_ascii=False, indent=1)
print("written", NEW, PNEW, MNEW)
