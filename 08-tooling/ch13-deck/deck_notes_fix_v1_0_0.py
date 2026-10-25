#!/usr/bin/env python3
"""deck_notes_fix v1.0.0 (chapter 13) - task TK_Deck13_ArcFix_Build: makes the NEW deck 2.1.0 and the NEW plan 2.1.0 from deck/plan 2.0.0 (both stay unchanged).
Only the speaker notes of slides 38 and 44 change (5 sentences each -> 3 sentences, same facts); the arc check v1_3_0 listed exactly these two. Usage: python3 -B deck_notes_fix_v1_0_0.py"""
__version__ = "1.0.0"
import json, os, shutil, subprocess, sys
from pptx import Presentation
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(os.path.dirname(HERE)); D = os.path.join(ROOT, "03-materials", "ch13")
OLD = os.path.join(D, "SEN0401_Ch13_BitcoinSecurity_3e_v2_0_0.pptx"); NEW = os.path.join(D, "SEN0401_Ch13_BitcoinSecurity_3e_v2_1_0.pptx")
NOTES = {38: "The secret is the height of a line at zero, and each share is its height at another point. Two points fix the line, so two shares rebuild the secret. Try other numbers and rerun; this is example 22 in the chapter code tab of the course page.",
         44: "Four figures from the story are recomputed here. If 850,000 coins were 7 percent of all coins, the supply was about 12.1 million, and a 36 percent cut turns 40 coins into 25.6. The sources disagree on the rest; this is example 25 in the chapter code tab of the course page."}
for p in (NEW, os.path.join(HERE, "deck_plan_v2_1_0.json")):
    if os.path.exists(p): sys.exit("REFUSED: %s exists" % p)
plan = json.load(open(os.path.join(HERE, "deck_plan_v2_0_0.json")))
for i, t in NOTES.items(): plan["slides"][i - 1]["notes"] = t
plan["meta"]["version"] = "2.1.0"; plan["meta"]["sub"] = "SEN0401, third-edition deck 2.1.0"
plan["meta"]["revision"] = "2026-10-25, task TK_Deck13_ArcFix_Build: the notes of slides 38 and 44 shortened to 3 sentences; nothing else changed"
json.dump(plan, open(os.path.join(HERE, "deck_plan_v2_1_0.json"), "w"), ensure_ascii=False, indent=1)
shutil.copy(OLD, NEW); pr = Presentation(NEW)
for i, t in NOTES.items(): pr.slides[i - 1].notes_slide.notes_text_frame.text = t
pr.save(NEW)
subprocess.check_call(["python3", os.path.join(os.path.dirname(HERE), "office_stamp_v1_0_0.py"), NEW, "2.1.0", "SEN0401 Chapter 13"])
