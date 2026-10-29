#!/usr/bin/env python3
"""Shared helpers of the chapter 6 deck v2.3.0 derivation (a NEW tool file): drawing in the deck's own palette with python-pptx, code colouring, and the corrected examples.
The deck v2.2.0 is read, never written. Units: 1 unit = 1/1000 of the slide width (the slide is 1000 x 562.5 units)."""
__version__ = "2.3.0"
import json, os, re
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

U = 9144
INK, MUTE, ORANGE, TEAL, CREAM, CARD, LINE, SOFT, WHITE, CODEBG = "2B2622", "6E655C", "D97A14", "2F6D62", "FAF6EF", "EFE9DC", "DFD6C7", "D7E4E0", "FFFFFF", "23272E"
KW = set("for in if else elif while def return import from as and or not is del with class try except lambda".split())
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(os.path.dirname(HERE))
MANIFEST_OLD = os.path.join(ROOT, "03-materials/ch06/page/chapter_code_ch06_v1_1_0.json")


def rgb(h): return RGBColor.from_string(h)


def one_line_compound_split(l):
    """('for ...:', 'body') when a compound statement keeps its body on the header line, else None (the colon is found outside brackets and strings)"""
    if not re.match(r"^\s*(for|while|if|elif|else|def|class|with|try|except|finally)\b", l): return None
    depth = 0; q = None
    for i, c in enumerate(l):
        if q:
            if c == q: q = None
        elif c in "'\"": q = c
        elif c in "([{": depth += 1
        elif c in ")]}": depth -= 1
        elif c == ":" and depth == 0 and l[i + 1:].strip(): return l[:i + 1], l[i + 1:].strip()
    return None


# The corrected examples: the 18 examples of the page's chapter code tab (manifest v1.1.0), with (1) no lambda (example 12 is rewritten with a for loop, same result),
# (2) every compound statement's body on its own indented line (examples 1, 6 and 17). Each original line maps to the first corrected line it became.
def corrected(code_lines):
    out = []; start = []
    for l in code_lines:
        start.append(len(out))
        sp = one_line_compound_split(l)
        if sp: out += [sp[0], "    " + sp[1]]
        else: out.append(l)
    return out, start


def fix_ex12(lines):
    """example 12: 'swap = lambda s: with_input_script(FIRST, s)' and the dict comprehension that used it become a plain for loop; the result is the same dict"""
    assert lines[1].startswith("swap = lambda") and lines[2].startswith("ids = {k: txid(swap(s))[:8]")
    return [lines[0], "ids = {}", "for k, s in enc.items():", "    ids[k] = txid(with_input_script(FIRST, s))[:8]"] + lines[3:], {0: 0, 1: 1, 2: 1, 3: 4, 4: 5}


def tokens(line):
    """(text, kind) pieces of one Python line: kw, str, num, com or plain"""
    out = []; i = 0
    pat = re.compile(r"(#.*$)|('[^']*'|\"[^\"]*\")|(\b0x[0-9a-fA-F]+\b|\b\d+(?:\.\d+)?\b)|(\b[A-Za-z_]\w*\b)")
    for m in pat.finditer(line):
        if m.start() > i: out.append((line[i:m.start()], "plain"))
        if m.group(1): out.append((m.group(1), "com"))
        elif m.group(2): out.append((m.group(2), "str"))
        elif m.group(3): out.append((m.group(3), "num"))
        else: out.append((m.group(4), "kw" if m.group(4) in KW else "plain"))
        i = m.end()
    if i < len(line): out.append((line[i:], "plain"))
    return out


CODE_COL = {"plain": "E6EDF3", "kw": "F2A860", "str": "9ECBFF", "num": "FFD28A", "com": "8B949E"}


def set_run(r, size, color, bold=False, italic=False, font="Calibri"):
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic; r.font.color.rgb = rgb(color); r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            from lxml import etree
            e = etree.SubElement(rPr, qn(tag))
        e.set("typeface", font)


def box(slide, name, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.08, lw=1.0):
    sp = slide.shapes.add_shape(shape, Emu(int(x * U)), Emu(int(y * U)), Emu(int(w * U)), Emu(int(h * U)))
    sp.name = name
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try: sp.adjustments[0] = adj
        except Exception: pass
    if fill: sp.fill.solid(); sp.fill.fore_color.rgb = rgb(fill)
    else: sp.fill.background()
    if line: sp.line.color.rgb = rgb(line); sp.line.width = Pt(lw)
    else: sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def text(slide, name, x, y, w, h, paras, anchor="t", wrap=True, inset=0, shape=None):
    """paras: list of paragraphs; a paragraph is a list of runs (text, size, color, bold, italic, font) or a dict {runs:[...], align:'l'|'c'|'r', space:pts}; shape: write into an existing shape instead of a new text box"""
    if shape is None:
        tb = slide.shapes.add_textbox(Emu(int(x * U)), Emu(int(y * U)), Emu(int(w * U)), Emu(int(h * U))); tb.name = name
    else:
        tb = shape
    tf = tb.text_frame; tf.word_wrap = wrap
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"): setattr(tf, side, Emu(int(inset * U)))
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    first = True
    for p in paras:
        d = p if isinstance(p, dict) else {"runs": p}
        para = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[d.get("align", "l")]
        if d.get("space"): para.space_after = Pt(d["space"])
        if d.get("line"): para.line_spacing = d["line"]
        for rr in d["runs"]:
            t, size, color = rr[0], rr[1], rr[2]
            bold = rr[3] if len(rr) > 3 else False; italic = rr[4] if len(rr) > 4 else False; font = rr[5] if len(rr) > 5 else "Calibri"
            r = para.add_run(); r.text = t; set_run(r, size, color, bold, italic, font)
    return tb


def code_paras(lines, size):
    return [[(t, size, CODE_COL[k], k == "kw", False, "Courier New") for t, k in tokens(l)] or [("", size, CODE_COL["plain"], False, False, "Courier New")] for l in lines]


def delete(shape):
    shape._element.getparent().remove(shape._element)


def est_lines(s, size, width_units, factor=0.5):
    """lines a text needs: an average character of a Calibri size s is about factor*s points wide; width in units (1 unit = 0.72 pt)"""
    per = max(1, int(width_units * 0.72 / (size * factor)))
    return sum(max(1, -(-len(p) // per)) for p in s.split("\n"))
