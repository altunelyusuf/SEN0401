#!/usr/bin/env python3
"""deck_plan_make v2.3.0 (a NEW tool file): writes deck_plan_v2_3_0.json, the plan that describes the chapter 6 deck v2.3.0 for the arc check (which reads plan and file together). It is DERIVED from the plan v2_1_0 (read, never written)
and from the deck v2.3.0 itself: the slides of v2.1.0 keep their kind, items, takeaway and background; the nine picture slides are added as content slides with a takeaway; the story slides keep their fact row (who, where, when, link: the
same facts now stand in the story sentences) and picture but lose the label beats; the notes of every slide are read from the deck; the slide numbers of the analogy list are mapped to the new order.
Usage: deck_plan_make_v2_3_0.py"""
import copy, json, os, sys
sys.dont_write_bytecode = True
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
from pptx import Presentation
import deck_enhance_v2_3_0 as E
ROOT = os.path.dirname(os.path.dirname(HERE))
old = json.load(open(os.path.join(HERE, "deck_plan_v2_1_0.json")))
deck = Presentation(os.path.join(ROOT, "03-materials/ch06/SEN0401_Ch06_Transactions_3e_v2_3_0.pptx"))
new_titles = {p[2]: p for p in E.PICS}
slides = []; oldpos = {}; k = 0
for i, s in enumerate(deck.slides, 1):
    title = next((sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.left // 9144 == 125 and sh.top // 9144 == 26), "")
    notes = s.notes_slide.notes_text_frame.text
    if title in new_titles:
        p = new_titles[title]
        slides.append({"kind": "content", "title": title, "sub": p[3], "items": [{"t": "picture", "native": True}], "notes": notes, "bg": "light", "take": p[4]})
    else:
        sl = copy.deepcopy(old["slides"][k]); oldpos[k + 1] = i; k += 1
        sl["notes"] = notes
        if any(it.get("t") == "factrow" for it in sl["items"]):
            sl["items"] = [it for it in sl["items"] if it.get("t") != "beats"]
        slides.append(sl)
assert k == 62, k
meta = copy.deepcopy(old["meta"]); meta["version"] = "2.3.0"; meta["total"] = len(slides)
meta["sub"] = "SEN0401, third-edition urgent deck: code corrected, concepts drawn, stories told"
for a in meta["analogies"]: a["slide"] = oldpos[a["slide"]]
json.dump({"meta": meta, "slides": slides}, open(os.path.join(HERE, "deck_plan_v2_3_0.json"), "x"), indent=1, ensure_ascii=False)
print("written deck_plan_v2_3_0.json:", len(slides), "slides")
