#!/usr/bin/env python3
"""sen0401_deck01_control_check v1.0.0 - the control's OWN checks of the chapter 1 deck v3_2_0 (story ST_Deck01_Charts, criterion AC9_Deck01_Charts). A NEW tool; it edits nothing.
Usage: sen0401_deck01_control_check_v1_0_0.py KIND
  sha       the deck v3_1_0 has the same SHA-256 in the commit before the executor's work (parent of the first commit) as now, and no later commit touches it
  parts     same zip part list; only docProps/core.xml differs (so every slide, note and chart part is byte-identical)
  textdiff  the text of every shape and table and the notes of every slide are identical
  links     the set of hyperlink addresses of v3_1_0 equals that of v3_2_0 (count printed)
  chartvals the control's OWN re-execution of the recorded chart expression of the supply chart (recording examples_out_v3_1_0.json, key _charts.ex04): its 8 values equal, as numbers rounded to 4 decimals, the plotted values of the chart part of the deck, and the chart is the one whose labels are E1..E8
  dcheck    the deck check v1_1_0 under venv314 on v3_2_0: exit 0, '0 mismatch(es)' and 'charts compared with 5 recorded expression(s)'
  arc0      the arc check v1_3_0 on v3_2_0 refuses nothing (exit 0)
Exit 0 only when the kind holds."""
__version__ = "1.0.0"
import hashlib, itertools, os, re, shutil, subprocess, sys, tempfile, zipfile
R = os.path.dirname(os.path.dirname(os.path.abspath(__file__))); D = os.path.join(R, "03-materials/ch01")
OLD = os.path.join(D, "SEN0401_Ch01_Introduction_3e_v3_1_0.pptx"); NEW = os.path.join(D, "SEN0401_Ch01_Introduction_3e_v3_2_0.pptx")
PARENT = "6f1f932"; V = "/home/claude/venv314/bin/python"; KEEP = os.path.join(R, "08-tooling/verify/iter32/bridge_dk01")
def sha(b): return hashlib.sha256(b).hexdigest()
def git(*a): return subprocess.run(["git", "-C", R] + list(a), capture_output=True).stdout
def done(ok, msg): print(msg); print("RESULT " + ("PASS" if ok else "FAIL")); sys.exit(0 if ok else 1)
def texts(path):
    from pptx import Presentation
    out = []
    for s in Presentation(path).slides:
        l = []
        for sh in s.shapes:
            if sh.has_text_frame: l += [p.text for p in sh.text_frame.paragraphs]
            if sh.has_table: l += [" | ".join(c.text for c in r.cells) for r in sh.table.rows]
        out.append((l, s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""))
    return out
def links(p):
    z = zipfile.ZipFile(p); o = set()
    for n in z.namelist():
        if re.match(r"ppt/slides/_rels/slide\d+\.xml\.rels$", n):
            for m in re.finditer(r"<Relationship [^>]*>", z.read(n).decode()):
                if "/hyperlink" in m.group(0): o.add(re.search(r'Target="([^"]+)"', m.group(0)).group(1))
    return o
def k_sha():
    rel = "03-materials/ch01/SEN0401_Ch01_Introduction_3e_v3_1_0.pptx"; b = git("show", "%s:%s" % (PARENT, rel)); n = open(os.path.join(R, rel), "rb").read(); ok = len(b) > 0 and sha(b) == sha(n)
    print("%s before(%s) %s after %s %s" % (rel, PARENT, sha(b), sha(n), "equal" if ok else "DIFFERENT")); ch = git("log", "--format=%h", "%s..HEAD" % PARENT, "--", rel).split(); print("commits since touching it: %s" % ch); done(ok and not ch, "sha")
def k_parts():
    a, b = zipfile.ZipFile(OLD), zipfile.ZipFile(NEW); na, nb = sorted(a.namelist()), sorted(b.namelist()); diff = sorted(n for n in na if n in nb and a.read(n) != b.read(n))
    print("part lists equal: %s (%d parts); differing parts: %s" % (na == nb, len(na), diff)); done(na == nb and diff == ["docProps/core.xml"], "parts")
def k_text():
    a, b = texts(OLD), texts(NEW); ch = [i + 1 for i in range(min(len(a), len(b))) if a[i] != b[i]]; print("slides %d and %d; slides whose text or notes differ: %s" % (len(a), len(b), ch)); done(len(a) == len(b) and not ch, "textdiff")
def k_links():
    a, b = links(OLD), links(NEW); print("distinct addresses v3_1_0: %d, v3_2_0: %d; only in old: %s; only in new: %s" % (len(a), len(b), sorted(a - b), sorted(b - a))); done(a == b and len(b) > 0, "links")
def k_chartvals():
    import json
    rec = json.load(open(os.path.join(R, "08-tooling/ch01-deck/examples_out_v3_1_0.json")))["_charts"]["ex04"]; sys.path.insert(0, os.path.join(R, "08-tooling/ch01-deck")); sys.path.insert(0, os.path.join(R, "08-tooling"))
    ns = {"HERE": os.path.join(R, "08-tooling/ch01-deck")}; exec(rec["setup"], ns); got = eval(rec["expr"], ns); z = zipfile.ZipFile(NEW); hit = None
    for n in sorted(x for x in z.namelist() if re.match(r"ppt/charts/chart\d+\.xml$", x)):
        t = z.read(n).decode(); cats = re.findall(r"<c:cat>.*?</c:cat>", t, re.S); lab = re.findall(r"<c:v>([^<]*)</c:v>", cats[0]) if cats else []
        if lab == ["E%d" % k for k in range(1, 9)]:
            vals = re.findall(r"<c:v>([^<]*)</c:v>", re.search(r"<c:val>.*?</c:val>", t, re.S).group(0)); hit = (n, [float(v) for v in vals])
    print("recorded expression re-executed: %s" % [(k, v) for k, v in got]); print("chart part %s plotted: %s" % (hit[0] if hit else None, hit[1] if hit else None))
    ok = hit is not None and len(got) == 8 and all(round(float(v), 4) == round(p, 4) for (k, v), p in zip(got, hit[1]))
    print("recorded _charts keys: %s" % sorted(json.load(open(os.path.join(R, "08-tooling/ch01-deck/examples_out_v3_1_0.json")))["_charts"])); done(ok, "chartvals")
def k_dcheck():
    r = subprocess.run([V, "-B", os.path.join(R, "08-tooling/sen0401_deck_check_v1_1_0.py"), "01", NEW], capture_output=True, text=True); out = (r.stdout + r.stderr).strip(); print(out[-700:]); print("exit %d" % r.returncode); done(r.returncode == 0 and "0 mismatch(es)" in out and "5 recorded expression(s)" in out, "dcheck")
def k_arc():
    r = subprocess.run([sys.executable, "-B", os.path.join(R, "08-tooling/sen0401_deck_arc_check_v1_3_0.py"), "01"], capture_output=True, text=True, env=dict(os.environ, SEN0401_DECK=NEW)); out = (r.stdout + r.stderr).strip(); print(out[-500:]); print("exit %d" % r.returncode); done(r.returncode == 0 and "REFUSED" not in out, "arc0")
if __name__ == "__main__":
    {"sha": k_sha, "parts": k_parts, "textdiff": k_text, "links": k_links, "chartvals": k_chartvals, "dcheck": k_dcheck, "arc0": k_arc}[sys.argv[1]]()
