#!/usr/bin/env python3
"""lecture_make v1_3_0 (SEN0401 chapter 5 page, task TK_Remedy_Page_Ch05_Build). The deck of chapter 5 is now v3.1.0 (68 slides, built with python-pptx; its plan file carries no READ cards or
link lists), so the deck's lecture export is read from the deck file itself (python-pptx 1.0.2): per slide the title, the speaker notes (say), the one-line takeaway (the shape named
"takeaway"), the READ card of a story slide (the hyperlink of the shape "live link") and the grouped links of the resources slide (the hyperlinks of the shapes named "resources <group>").
Writes, as NEW files (nothing committed is edited):
 (1) lecture_v1_3_0.json - the deck export (the form of lecture_v1_1_0.json: n, kind, title, say, concept = the takeaway, read, links);
 (2) lecture_v1_3_1.json - as pages_deck_links_make_v1_2_0 made lecture_v1_2_1.json: the takeaway moves to "takeaway", and a story slide gets as "concept" the first concept of its story that is a concept of the page;
 (3) resources_v1_3_0.json - from resources_v1_2_0.json as that tool made it: an entry whose address is a link of the deck's resources slide takes the deck's label (and its line of what the reader finds, taken from the
     earlier export or list when the same address was there); a link of the deck that the list does not have is added.
Usage: lecture_make_v1_3_0.py"""
__version__ = "1.3.0"
import glob, importlib, json, os, re, sys
from urllib.parse import urlparse
from pptx import Presentation
H = os.path.dirname(os.path.abspath(__file__)); T = os.path.dirname(H); REPO = os.path.dirname(T); sys.path.insert(0, T)
vk = lambda p: [int(x) for x in re.findall(r"_v(\d+)_(\d+)_(\d+)", p)[-1]]
DECK = "SEN0401_Ch05_WalletRecovery_3e_v3_1_0.pptx"
prs = Presentation(os.path.join(REPO, "03-materials", "ch05", DECK)); N = len(prs.slides)
prev = json.load(open(os.path.join(H, "lecture_v1_2_0.json"))); old_finds = {l["url"]: l.get("finds", "") for s in prev["slides"] for l in s.get("links", [])}
res = json.load(open(os.path.join(H, "resources_v1_2_0.json"))); old_finds.update({r["url"]: r.get("why", "") for r in res["links"]})
slides = []
for i, sl in enumerate(prs.slides, 1):
    d = {"n": i, "kind": "title" if i == 1 else "closing" if i == N else "content", "title": sl.shapes.title.text_frame.text.strip(), "say": sl.notes_slide.notes_text_frame.text.strip(), "concept": ""}
    groups = []
    for sh in sl.shapes:
        if sh.name == "takeaway" and sh.has_text_frame: d["concept"] = sh.text_frame.text.strip()
        if sh.name == "live link" and sh.has_text_frame:
            r = next(r for p in sh.text_frame.paragraphs for r in p.runs if r.hyperlink.address); d["read"] = {"label": r.text.strip(), "url": r.hyperlink.address}
        if sh.name.startswith("resources ") and sh.has_text_frame:
            g = sh.name[len("resources "):]
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.hyperlink.address: groups.append({"group": g, "label": r.text.strip(), "url": r.hyperlink.address, "finds": old_finds.get(r.hyperlink.address, "")})
    if groups: d["links"] = groups
    slides.append(d)
raw = {"chapter": 5, "deck": DECK, "slides": slides}
json.dump(raw, open(os.path.join(H, "lecture_v1_3_0.json"), "w"), indent=1, ensure_ascii=False)
# (2) the page form
pd = json.load(open(sorted(glob.glob(os.path.join(H, "page_data_v*.json")), key=vk)[-1])); nodes = {x["id"]: x for x in pd["nodes"]}
mod = importlib.import_module(os.path.basename(sorted(glob.glob(os.path.join(T, "sen0401_ch05_stories_v*.py")), key=vk)[-1])[:-3])
lec = json.loads(json.dumps(raw)); k = 0
for s in lec["slides"]:
    s["takeaway"] = s.pop("concept", "")
    if s.get("read"):
        st = next((x for x in mod.STORIES if x["title"] == s["title"] or x.get("link") == s["read"]["url"]), None)
        if st is None: raise SystemExit("slide %d shows a story the stories module does not have: %r" % (s["n"], s["title"]))
        c = next((c for c in st["concepts"] if c in nodes), None); assert c, st["id"]; s["concept"] = c; k += 1
lec["_note"] = "lecture_v1_3_0.json (the export of the chapter 5 deck v3.1.0, read from the deck file) with the deck's takeaway as takeaway and, on each story slide, the first concept of the story that the page has as concept; made by lecture_make_v1_3_0.py"
json.dump(lec, open(os.path.join(H, "lecture_v1_3_1.json"), "w"), indent=1, ensure_ascii=False)
# (3) the resources
new = [l for s in raw["slides"] for l in s.get("links", [])]; nu = {l["url"]: l for l in new}
KIND = lambda g, u: "video" if "youtube" in u else "standard" if re.search(r"standard|bip", g, re.I) or "bips.dev" in u else "seminal" if re.search(r"seminal|papers", g, re.I) else "community" if re.search(r"communit", g, re.I) else "documentation"
root = pd["nodes"][0]["id"]
def concept_for(lbl):
    w = [x for x in re.findall(r"[A-Za-z]{4,}", lbl.lower()) if x not in ("bitcoin", "guide", "news", "spec", "video", "text")]
    for nid, n in nodes.items():
        if any(x in n["label"].lower() for x in w): return nid
    return root
out = []; relabel = added = 0; have = set()
for r in res["links"]:
    r = dict(r)
    if r["url"] in nu: d = nu[r["url"]]; r["title"] = d["label"]; r["why"] = d.get("finds") or r["why"]; relabel += 1
    have.add(r["url"]); out.append(r)
for l in new:
    if l["url"] in have: continue
    out.append({"kind": KIND(l["group"], l["url"]), "title": l["label"], "url": l["url"], "source": urlparse(l["url"]).netloc.replace("www.", ""), "why": l.get("finds") or "Listed on the lecture's resources slide, group %s." % l["group"], "concept": concept_for(l["label"])}); have.add(l["url"]); added += 1
res["links"] = out; res["_version"] = "1.3.0"
res["_note"] = (res.get("_note", "") + " 1.3.0 (after the deck v3.1.0, from lecture_v1_3_0.json): %d entries carry the deck's label, %d deck links added that the list did not have." % (relabel, added))
json.dump(res, open(os.path.join(H, "resources_v1_3_0.json"), "w"), indent=1, ensure_ascii=False)
print("lecture_v1_3_0.json: %d slides, %d READ cards, %d resource links; lecture_v1_3_1.json: %d story slides with a concept; resources_v1_3_0.json: %d entries (relabelled %d, added %d)" % (N, sum(1 for s in raw["slides"] if s.get("read")), len(new), k, len(out), relabel, added))
