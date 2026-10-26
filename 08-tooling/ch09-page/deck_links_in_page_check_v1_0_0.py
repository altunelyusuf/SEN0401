#!/usr/bin/env python3
"""deck_links_in_page_check v1_0_0 (executor's own listing, not a gate): every hyperlink of the chapter 9 deck (runs, shapes and tables) must appear in the page text. Usage: deck_links_in_page_check_v1_0_0.py DECK PAGE"""
import sys
from pptx import Presentation
p = Presentation(sys.argv[1]); urls = set()
def runs(tf):
    for para in tf.paragraphs:
        for r in para.runs:
            if r.hyperlink and r.hyperlink.address: urls.add(r.hyperlink.address)
def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6: walk(sh.shapes)
        if sh.has_text_frame: runs(sh.text_frame)
        try:
            if sh.click_action and sh.click_action.hyperlink.address: urls.add(sh.click_action.hyperlink.address)
        except Exception: pass
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for c in row.cells: runs(c.text_frame)
for s in p.slides: walk(s.shapes)
page = open(sys.argv[2], encoding="utf8").read()
miss = [u for u in sorted(urls) if u not in page and u.replace("&", "&amp;") not in page]
print("%d deck hyperlinks; not in the page: %d" % (len(urls), len(miss)))
for u in miss: print("  MISSING", u)
sys.exit(1 if miss else 0)
