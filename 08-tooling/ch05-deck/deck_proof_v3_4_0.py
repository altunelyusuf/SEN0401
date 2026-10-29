#!/usr/bin/env python3
"""Proof of the deck v3.4.0 against v3.3.0 (a NEW tool file): slide text and notes diff, differing package parts, contrast reading of the closing slide. Usage: deck_proof_v3_4_0.py"""
import os, re, sys, zipfile, difflib
from pptx import Presentation
R = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
A = R + "/03-materials/ch05/SEN0401_Ch05_WalletRecovery_3e_v3_3_0.pptx"; B = R + "/03-materials/ch05/SEN0401_Ch05_WalletRecovery_3e_v3_4_0.pptx"
def dump(f):
    out = []
    for i, s in enumerate(Presentation(f).slides, 1):
        for sh in s.shapes:
            if sh.has_text_frame: out += ["s%d|%s" % (i, p.text) for p in sh.text_frame.paragraphs]
        if s.has_notes_slide: out.append("s%d|NOTES|%s" % (i, s.notes_slide.notes_text_frame.text))
    return out
a, b = dump(A), dump(B)
d = [l for l in difflib.unified_diff(a, b, lineterm="", n=0) if not l.startswith(("---", "+++", "@@"))]
print("slide text and notes: %d lines in v3_3_0, %d in v3_4_0, %d difference line(s)" % (len(a), len(b), len(d)))
print("NO DIFFERENCE in any slide text or any speaker note" if not d else "DIFFERENCE: " + str(d[:10]))
za, zb = zipfile.ZipFile(A), zipfile.ZipFile(B)
na, nb = za.namelist(), zb.namelist()
print("package parts only in v3_3_0:", sorted(set(na) - set(nb)), "only in v3_4_0:", sorted(set(nb) - set(na)))
print("same part order:", na == nb)
print("parts with different bytes:", sorted(n for n in na if n in nb and za.read(n) != zb.read(n)))
def lum(h):
    c = [int(h[i:i+2], 16) / 255 for i in (0, 2, 4)]
    c = [x / 12.92 if x <= 0.03928 else ((x + 0.055) / 1.055) ** 2.4 for x in c]
    return 0.2126 * c[0] + 0.7152 * c[1] + 0.0722 * c[2]
def cr(f, g): x, y = sorted((lum(f), lum(g)), reverse=True); return (x + 0.05) / (y + 0.05)
x = zb.read("ppt/slides/slide68.xml").decode()
bg = re.search(r'<p:bg>.*?val="([0-9A-F]{6})"', x).group(1)
print("closing slide background %s, relative luminance %.3f (dark below 0.5)" % (bg, lum(bg)))
low = 0
for sp in re.findall(r"<p:sp>.*?</p:sp>", x):
    name = re.search(r'name="([^"]*)"', sp).group(1)
    fill = re.search(r"</a:prstGeom><a:solidFill><a:srgbClr val=\"([0-9A-F]{6})\"", sp)
    under = fill.group(1) if fill else bg
    for rp, fg, t in re.findall(r'<a:rPr([^>]*)><a:solidFill><a:srgbClr val="([0-9A-F]{6})"/></a:solidFill>.*?<a:t>([^<]*)</a:t>', sp):
        pt = int(re.search(r'sz="(\d+)"', rp).group(1)) / 100; bold = 'b="1"' in rp
        large = pt >= 18 or (pt >= 14 and bold); need = 3.0 if large else 4.5
        r = cr(fg, under); ok = r >= need; low += not ok
        print("  %-12s %4.1fpt%s text %s on %s contrast %.2f:1 needs %.1f %s | %s" % (name, pt, " bold" if bold else "", fg, under, r, need, "OK" if ok else "BELOW", t[:50]))
print("CONTRAST RESULT:", "every text run meets WCAG AA (4.5:1 normal text, 3:1 large text)" if not low else "%d run(s) below WCAG AA" % low)
