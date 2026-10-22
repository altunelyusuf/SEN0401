#!/usr/bin/env python3
"""BUILDER'S OWN MEASUREMENT of the chapter 5 deck against the deck standard D1-D12 (python-pptx 1.0.2).
This is NOT the independent deck standard check (TK_Remedy_DeckCheck): the builder never judges its own work with a tool it wrote;
this script only lets the builder see its numbers while building. Usage: deck_selfmeasure_ch05_v1_0_0.py <deck.pptx> [manifest.json]"""
__version__ = "1.0.0"
import json, os, re, sys
from pptx import Presentation
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
deck = sys.argv[1]
man = json.load(open(sys.argv[2])) if len(sys.argv) > 2 else None
prs = Presentation(deck)
slides = list(prs.slides)
N = len(slides)
res = {}


def words(t):
    return len(t.split())


def texts(sl):
    out = []
    for sh in sl.shapes:
        out += _texts(sh)
    return out


def _texts(sh):
    out = []
    if sh.shape_type == 6:                                  # group
        for g in sh.shapes:
            out += _texts(g)
    elif sh.has_text_frame and sh.text_frame.text.strip():
        out.append((sh, sh.text_frame.text))
    elif getattr(sh, "has_table", False) and sh.has_table:
        for r in sh.table.rows:
            for c in r.cells:
                if c.text_frame.text.strip():
                    out.append((sh, c.text_frame.text))
    return out


# D1
bad = [i for i, sl in enumerate(slides, 1) if not any(re.fullmatch(r"%d / %d" % (i, N), t.strip()) for _, t in texts(sl))]
res["D1"] = (not bad, "%d slides numbered n / %d; missing %s" % (N - len(bad), N, bad))
# D2
mx, worst, over = 0, None, []
mxcell = 0
for i, sl in enumerate(slides, 1):
    for sh, t in texts(sl):
        w = words(t)
        if sh.has_text_frame if hasattr(sh, "has_text_frame") else False:
            if w > mx:
                mx, worst = w, (i, t[:50])
            if w > 30:
                over.append((i, w, t[:40]))
        else:
            mxcell = max(mxcell, w)
res["D2"] = (not over, "longest text box %d words (slide %s); over 30: %s; longest table cell %d words" % (mx, worst, over, mxcell))
# D3
code_slides, d3bad = [], []
for i, sl in enumerate(slides, 1):
    codes = [sh for sh in sl.shapes if sh.name == "code"]
    if not codes:
        continue
    code_slides.append(i)
    expl = [sh for sh in sl.shapes if sh.name == "explanation"]
    vis = [sh for sh in sl.shapes if sh.has_chart or sh.has_table or (sh.shape_type == 6 and sh.name.startswith("visual"))] if True else []
    vis = [sh for sh in sl.shapes if getattr(sh, "has_chart", False) and sh.has_chart or getattr(sh, "has_table", False) and sh.has_table or (sh.shape_type == 6 and sh.name.startswith("visual"))]
    if len(expl) != len(codes) or not vis:
        d3bad.append((i, len(codes), len(expl), len(vis)))
res["D3"] = (not d3bad and bool(code_slides), "%d code slides; each has an explanation beside every code group and a result visual; failing %s" % (len(code_slides), d3bad))
# D4 analogies for the key terms
ttl = open(os.path.join(ROOT, "03-materials", "ch05", "rdodi", "sen0401_ch05_research_v1_1_1.ttl"), encoding="utf-8").read()
conc = re.findall(r'rdfs:label "([^"]+)"@en ; sen0401:conceptKind "(Section|Concept)"', ttl)
terms = [l for l, k in conc if k == "Concept"]
secs = [l for l, k in conc if k == "Section"]
an_first_cols = set()
for sl in slides:
    if sl.shapes.title.text_frame.text.startswith("Everyday pictures"):
        for sh in sl.shapes:
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in list(sh.table.rows)[1:]:
                    an_first_cols.add(r.cells[0].text_frame.text.strip().lower())
keyc = terms + secs
miss = [t for t in keyc if t.lower() not in an_first_cols]
res["D4"] = (not miss, "%d of %d research concepts (44 named terms + 26 book sections) have an everyday analogy row; missing %s" % (len(keyc) - len(miss), len(keyc), miss))
# D5
sl5, d5bad = [], []
for i, sl in enumerate(slides, 1):
    ts = [t for _, t in texts(sl)]
    if all(l in ts for l in ("WHO", "WHAT", "WHEN", "WHERE", "WHY", "HOW")):
        sl5.append(i)
        pic = any(sh.shape_type == 13 for sh in sl.shapes)
        link = any(r.hyperlink.address for sh in sl.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs)
        if not (pic and link):
            d5bad.append(i)
res["D5"] = (len(sl5) >= 3 and not d5bad, "%d story slides %s with who/what/when/where/why/how, a picture and a live link; failing %s" % (len(sl5), sl5, d5bad))
# D6
charts = sum(1 for sl in slides for sh in sl.shapes if getattr(sh, "has_chart", False) and sh.has_chart)
pics = [(i, sh.name) for i, sl in enumerate(slides, 1) for sh in sl.shapes if sh.shape_type == 13]
res["D6"] = (charts > 0, "%d native chart objects; %d pictures (book figures and story photographs)" % (charts, len(pics)))
# D7
nw = [(i, words(sl.notes_slide.notes_text_frame.text)) for i, sl in enumerate(slides, 1)]
res["D7"] = (all(w <= 60 for _, w in nw), "longest note %d words; over 60: %s" % (max(w for _, w in nw), [x for x in nw if x[1] > 60]))
# D8
r8 = [i for i, sl in enumerate(slides, 1) if sl.shapes.title.text_frame.text.startswith("Where to go")]
nl = sum(1 for sh in slides[r8[0] - 1].shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs if r.hyperlink.address) if r8 else 0
res["D8"] = (bool(r8) and nl >= 8, "resources slide %s with %d links" % (r8, nl))
# D9
bgs = set()
for sl in slides:
    f = sl.background.fill
    bgs.add(str(f.fore_color.rgb) if f.type == 1 else "none")
res["D9"] = (len(bgs) == 1, "backgrounds used: %s" % sorted(bgs))
# D10
dv = [i for i, sl in enumerate(slides, 1) if re.match(r"Part \d+:", sl.shapes.title.text_frame.text)]  # dividers carry no brand chip
d10bad = []
for i in dv:
    sl = slides[i - 1]
    rows = [t for sh, t in texts(sl) if sh.name == "concept meaning" and len(t.split()) >= 5]
    vis = [sh for sh in sl.shapes if getattr(sh, "has_table", False) and sh.has_table]
    if len(rows) < 4 or not vis:
        d10bad.append(i)
res["D10"] = (bool(dv) and not d10bad, "%d dividers %s, each with named concepts, a line of meaning each, and a native table; failing %s" % (len(dv), dv, d10bad))
# D11
alltext = " ".join(re.sub(r"\s+", " ", t.lower()) for sl in slides for _, t in texts(sl))
for sl in slides:
    for sh in sl.shapes:
        if getattr(sh, "has_chart", False) and sh.has_chart:
            alltext += " " + " ".join(str(c).lower() for c in sh.chart.plots[0].categories)
alllabels = terms + secs
missed = [l for l in alllabels if re.sub(r"\s+", " ", l.lower()) not in alltext]
share = 100.0 * (len(alllabels) - len(missed)) / len(alllabels)
res["D11"] = (share >= 90, "%d of %d research concepts appear on a slide = %.1f percent; not found: %s" % (len(alllabels) - len(missed), len(alllabels), share, missed))
# D12
d12bad = []
for i, sl in enumerate(slides, 1):
    for sh in sl.shapes:
        if sh.name == "code":
            tf = sh.text_frame
            fs = max(r.font.size.pt for p in tf.paragraphs for r in p.runs)
            ln = max(len("".join(r.text for r in p.runs)) for p in tf.paragraphs)
            need = ln * 0.6 * fs + 0.2 * 72
            have = Emu(sh.width).pt
            if tf.word_wrap or need > have:
                d12bad.append((i, ln, round(need), round(have)))
res["D12"] = (not d12bad, "wrap off on every code box and the longest line fits its box; failing %s" % d12bad)
# manifest order
if man:
    ok = [ex["slide"] for ex in man["examples"]] == code_slides
    res["D3b"] = (ok, "manifest lists %d examples; slide numbers equal the code slides' order: %s" % (len(man["examples"]), ok))

for k in sorted(res, key=lambda x: (int(re.match(r"D(\d+)", x).group(1)), x)):
    ok, msg = res[k]
    print("%s %s  %s" % (k, "PASS" if ok else "FAIL", msg))
sys.exit(0 if all(ok for ok, _ in res.values()) else 1)
