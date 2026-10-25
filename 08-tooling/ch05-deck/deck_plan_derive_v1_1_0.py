#!/usr/bin/env python3
"""deck_plan_derive v1.1.0 - derives deck_plan_v3_3_0.json (the schema sen0401_deck_arc_check_v1_3_0.py reads) from the delivered chapter 5 deck and the plan deck_plan_v3_1_0.json.
Reads only; writes the one new plan file. The deck is not edited or rebuilt. Usage: python3 -B deck_plan_derive_v1_1_0.py OLDPLAN.json OUT.json
v1.1.0 over v1.0.0: reads the deck v3_2_0; OLDPLAN.json is the old-schema plan its builder wrote (picture files, chart items with labels and series); the new plan KEEPS every item and
meta field of that old schema (img path, chart title/labels/series, ...) and ADDS the arc check's kind, bg, sub, take, the arc items (table, numlist, concepts, factrow, links, bullets) and meta.analogies."""
import json, os, re, sys
from pptx import Presentation
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
DECK = os.path.join(ROOT, "03-materials/ch05/SEN0401_Ch05_WalletRecovery_3e_v3_2_0.pptx")
OLD = sys.argv[1]
old = json.load(open(OLD)); prs = Presentation(DECK); assert len(old["slides"]) == len(prs.slides)
N = len(prs.slides)

def bg(slide):
    x = slide._element.xpath(".//p:bg//a:srgbClr/@val")
    if not x: return "light"
    r, g, b = (int(x[0][i:i + 2], 16) for i in (0, 2, 4))
    return "dark" if (0.299 * r + 0.587 * g + 0.114 * b) < 110 else "light"

slides, analogies = [], []
for i, sl in enumerate(prs.slides, 1):
    o = old["slides"][i - 1]; sh = list(sl.shapes); by = lambda n: [s for s in sh if s.name == n]
    txt = lambda n: (by(n)[0].text_frame.text if by(n) and by(n)[0].has_text_frame else "")
    title = sl.shapes.title.text_frame.text
    assert title == o["title"], (i, title, o["title"])
    notes = sl.notes_slide.notes_text_frame.text if sl.has_notes_slide else ""
    kind = "title" if i == 1 else "closing" if i == N else "section" if re.match(r"Part \d+:", title) else "content"
    items = []
    for it in o["items"]:
        it = dict(it)
        if it.get("t") == "img" and txt("credit"): it["cap"] = txt("credit")
        items.append(it)
    tabs = [s for s in sh if s.has_table]
    for t in tabs:
        rows = [[c.text for c in r.cells] for r in t.table.rows]
        items.append({"t": "table", "name": t.name, "rows": rows})
        if t.name == "analogy table":
            for r in rows[1:]: analogies.append({"concept": r[0], "analogy": r[1], "slide": i})
    if i == 2: items.append({"t": "numlist", "items": [txt("route row") if False else s.text_frame.text for s in by("route row")]})
    if i == 3: items.append({"t": "numlist", "items": [s.text_frame.text for s in by("question")]})
    if kind == "section": items.append({"t": "concepts", "rows": [[a.text_frame.text, b.text_frame.text] for a, b in zip(by("concept"), by("concept meaning"))]})
    if kind == "closing": items.append({"t": "bullets", "items": [s.text_frame.text for s in by("outcome")], "dark": bg(sl) == "dark"})
    if by("5N1K label"):
        lab = {l.text_frame.text: t.text_frame.text for l, t in zip(by("5N1K label"), by("5N1K text"))}
        ll = by("live link")[0]; run = [r for pa in ll.text_frame.paragraphs for r in pa.runs if r.hyperlink.address][0]
        items.append({"t": "factrow", "who": lab["WHO"], "what": lab["WHAT"], "when": lab["WHEN"], "where": lab["WHERE"], "why": lab["WHY"], "how": lab["HOW"],
                      "link": run.hyperlink.address, "linkText": run.text})
    rg = [s for s in sh if s.name.startswith("resources")]
    if rg:
        groups = []
        for g in rg:
            ps = g.text_frame.paragraphs; rows = []
            for pa in ps[1:]:
                for r in pa.runs:
                    if r.hyperlink.address: rows.append([r.text, r.hyperlink.address])
            groups.append({"head": ps[0].text, "rows": rows})
        items.append({"t": "links", "groups": groups})
    s = {"kind": kind, "title": title, "items": items, "notes": notes, "bg": bg(sl)}
    assert notes == o["notes"], i
    if txt("subtitle"): s["sub"] = txt("subtitle")
    if txt("takeaway"): s["take"] = txt("takeaway")
    slides.append(s)
meta = dict(old["meta"]); meta["version"] = "3.3.0"; meta["total"] = N
meta["schema"] = "arc check schema (kind, title, sub, take, bg, items, notes; meta.analogies); derived from " + os.path.basename(DECK) + " and the old-schema plan of its builder by deck_plan_derive_v1_1_0.py"
meta["analogies"] = analogies
json.dump({"meta": meta, "slides": slides}, open(sys.argv[2], "w"), indent=1, ensure_ascii=False)
print("wrote", sys.argv[2], N, "slides,", len(analogies), "analogies")
