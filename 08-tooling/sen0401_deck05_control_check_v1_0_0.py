#!/usr/bin/env python3
"""sen0401_deck05_control_check v1.0.0 - the control's OWN checks of the chapter 5 deck v3_3_0 (story ST_Deck05_Examples, criterion AC9_Deck05_Examples). A NEW tool; it edits nothing.
Usage: sen0401_deck05_control_check_v1_0_0.py KIND
  sha       the deck v3_2_0 has the same SHA-256 in the commit before the executor's work (parent of the first commit) as now, and no later commit touches it
  parts     the zip part lists of v3_2_0 and v3_3_0 are equal and only docProps/core.xml differs
  textdiff  the text of every shape and table and the notes of every slide are identical
  links     the set of hyperlink addresses of v3_2_0 equals that of v3_3_0 (count printed)
  charts    the chart parts of v3_3_0 equal the chart parts of v3_2_0 byte for byte, (the plan's chart items are compared by the dcheck kind) (compared by the deck check, kind dcheck)
  dcheck    the deck check v1_1_0 under venv314 on v3_3_0: exit 0 and '0 mismatch(es)'
  arcclose  the arc check v1_3_0 on v3_3_0: every refusal is the closing slide one (the pending owner decision CR3_GateArcDark), and the refusals of v3_2_0 are the same lines
Exit 0 only when the kind holds."""
__version__ = "1.0.0"
import hashlib, json, os, re, subprocess, sys, zipfile
R = os.path.dirname(os.path.dirname(os.path.abspath(__file__))); D = os.path.join(R, "03-materials/ch05")
OLD = os.path.join(D, "SEN0401_Ch05_WalletRecovery_3e_v3_2_0.pptx"); NEW = os.path.join(D, "SEN0401_Ch05_WalletRecovery_3e_v3_3_0.pptx")
PARENT = "af575d2"; V = "/home/claude/venv314/bin/python"
def sha(b): return hashlib.sha256(b).hexdigest()
def git(*a): return subprocess.run(["git", "-C", R] + list(a), capture_output=True).stdout
def done(ok, msg): print(msg); print("RESULT " + ("PASS" if ok else "FAIL")); sys.exit(0 if ok else 1)
def texts(path):
    from pptx import Presentation
    out = []
    for s in Presentation(path).slides:
        l = []
        for sh in s.shapes:
            if sh.has_text_frame: l += [p.text for p in sh.text_frame.paragraphs]
            if sh.has_table: l += [" | ".join(c.text for c in r.cells) for r in sh.table.rows]
        out.append((l, s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""))
    return out
def links(p):
    z = zipfile.ZipFile(p); o = set()
    for n in z.namelist():
        if re.match(r"ppt/slides/_rels/slide\d+\.xml\.rels$", n):
            for m in re.finditer(r"<Relationship [^>]*>", z.read(n).decode()):
                if "/hyperlink" in m.group(0): o.add(re.search(r'Target="([^"]+)"', m.group(0)).group(1))
    return o
def k_sha():
    rel = "03-materials/ch05/SEN0401_Ch05_WalletRecovery_3e_v3_2_0.pptx"; b = git("show", "%s:%s" % (PARENT, rel)); n = open(os.path.join(R, rel), "rb").read(); ok = len(b) > 0 and sha(b) == sha(n)
    print("%s before(%s) %s after %s %s" % (rel, PARENT, sha(b), sha(n), "equal" if ok else "DIFFERENT"))
    ch = git("log", "--format=%h", "%s..HEAD" % PARENT, "--", rel).split(); print("commits since touching it: %s" % ch); done(ok and not ch, "sha")
def k_parts():
    a, b = zipfile.ZipFile(OLD), zipfile.ZipFile(NEW); na, nb = sorted(a.namelist()), sorted(b.namelist()); diff = [n for n in na if n in nb and a.read(n) != b.read(n)]
    print("part lists equal: %s (%d parts); differing parts: %s" % (na == nb, len(na), diff)); done(na == nb and diff == ["docProps/core.xml"], "parts")
def k_text():
    a, b = texts(OLD), texts(NEW); ch = [i + 1 for i in range(min(len(a), len(b))) if a[i] != b[i]]; print("slides %d and %d; slides whose text or notes differ: %s" % (len(a), len(b), ch)); done(len(a) == len(b) and not ch, "textdiff")
def k_links():
    a, b = links(OLD), links(NEW); print("distinct addresses v3_2_0: %d, v3_3_0: %d; only in old: %s; only in new: %s" % (len(a), len(b), sorted(a - b), sorted(b - a))); done(a == b and len(b) > 0, "links")
def k_charts():
    a, b = zipfile.ZipFile(OLD), zipfile.ZipFile(NEW); ca = sorted(n for n in a.namelist() if re.match(r"ppt/charts/chart\d+\.xml$", n)); cb = sorted(n for n in b.namelist() if re.match(r"ppt/charts/chart\d+\.xml$", n))
    same = ca == cb and all(a.read(n) == b.read(n) for n in ca)
    print("chart parts v3_2_0: %d, v3_3_0: %d, byte-identical: %s (the comparison of the charts with the plan's chart items is made by the deck check, kind dcheck)" % (len(ca), len(cb), same)); done(same and len(cb) > 0, "charts")
def k_dcheck():
    r = subprocess.run([V, "-B", os.path.join(R, "08-tooling/sen0401_deck_check_v1_1_0.py"), "05", NEW], capture_output=True, text=True); out = (r.stdout + r.stderr).strip(); print(out[-900:]); print("exit %d" % r.returncode)
    done(r.returncode == 0 and "0 mismatch(es)" in out, "dcheck")
def k_arc():
    def ref(p):
        r = subprocess.run([sys.executable, "-B", os.path.join(R, "08-tooling/sen0401_deck_arc_check_v1_3_0.py"), "05"], capture_output=True, text=True, env=dict(os.environ, SEN0401_DECK=p)); return r.returncode, [l.strip() for l in (r.stdout + r.stderr).splitlines() if "REFUSED" in l]
    rn, ln = ref(NEW); ro, lo = ref(OLD); print("v3_3_0: exit %d, refusals %s" % (rn, ln)); print("v3_2_0: exit %d, refusals %s" % (ro, lo))
    done(ln == lo and all("closing" in l for l in ln), "arcclose")
if __name__ == "__main__":
    {"sha": k_sha, "parts": k_parts, "textdiff": k_text, "links": k_links, "charts": k_charts, "dcheck": k_dcheck, "arcclose": k_arc}[sys.argv[1]]()
