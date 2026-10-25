#!/usr/bin/env python3
"""sen0401_ch05fix2_control_check v1_0_0 - the lineage agent's OWN checks of the chapter 5 deck fix 2 (criterion AC9_Ch05DeckFix2). Usage: sen0401_ch05fix2_control_check_v1_0_0.py plan | arcone | codeslides
plan: the newest plan (v3_3_0) has as many slides as the newest deck (v3_2_0), every plan title occurs on the deck slide of the same number, the last slide is declared as it is (kind closing, bg light, equal to the deck's real background), and the plan keeps every slide-level key, every item key and every meta key that the plan v3_1_0 has (no field dropped).
arcone: the arc check v1_3_0 on chapter 5 prints exactly ONE refusal, 'the last slide is not the dark closing slide' (the recorded gate contradiction ex:CR3_GateArcDark), and no other; the last slide is NOT dark (no switch).
codeslides: the slides of the chapter-code manifest examples equal the slides of the newest deck that hold a '>>> ' paragraph, in order, and each example's first code line appears on its slide."""
import glob, json, os, re, subprocess, sys
from pptx import Presentation
T = os.path.dirname(os.path.abspath(__file__)); R = os.path.dirname(T); vk = lambda p: [int(x) for x in re.findall(r"_v(\d+)_(\d+)_(\d+)", p)[-1]]
deck = sorted(glob.glob(R + "/03-materials/ch05/*.pptx"), key=vk)[-1]; plan = sorted(glob.glob(T + "/ch05-deck/deck_plan_v*.json"), key=vk)[-1]; old = T + "/ch05-deck/deck_plan_v3_1_0.json"
def keys(P):
    s, i = set(), set()
    for sl in P["slides"]:
        s |= set(sl)
        for it in sl.get("items", []) or []:
            if isinstance(it, dict): i |= set(it)
    return s, i, set(P.get("meta", {}))
def planc():
    P = json.load(open(plan)); O = json.load(open(old)); D = list(Presentation(deck).slides); print(os.path.basename(plan), len(P["slides"]), "slides;", os.path.basename(deck), len(D), "slides"); norm = lambda s: re.sub(r"\W+", " ", s.lower()).strip()
    bad = [i for i, (p, d) in enumerate(zip(P["slides"], D), 1) if norm(p["title"]) not in norm(" ".join(sh.text_frame.text for sh in d.shapes if sh.has_text_frame))]; print("plan titles not on the deck slide of the same number:", bad)
    last = P["slides"][-1]; fill = D[-1].background.fill; rgb = str(fill.fore_color.rgb) if fill.type is not None else ""; print("last slide plan:", last.get("kind"), last.get("bg"), "| deck background", rgb)
    ps, pi, pm = keys(P); os_, oi, om = keys(O); miss = (os_ - ps, oi - pi, om - pm); print("fields of v3_1_0 missing in", os.path.basename(plan), ": slide keys", sorted(miss[0]), "item keys", sorted(miss[1]), "meta keys", sorted(miss[2]))
    return len(P["slides"]) == len(D) and not bad and last.get("kind") == "closing" and last.get("bg") == "light" and rgb.upper() == "FAF6EF" and not any(miss)
def arcone():
    r = subprocess.run([sys.executable, "-B", T + "/sen0401_deck_arc_check_v1_3_0.py", "05"], capture_output=True, text=True); print(r.stdout.strip()); ref = [l.strip() for l in r.stdout.splitlines() if l.strip().startswith("REFUSED")]; print("arc check exit", r.returncode, "refusals", len(ref))
    return len(ref) == 1 and "the last slide is not the dark closing slide" in ref[0] and r.returncode == 1
def codeslides():
    m = json.load(open(sorted(glob.glob(R + "/03-materials/ch05/page/chapter_code_ch05_v*.json"), key=vk)[-1])); ex = m["examples"]; D = list(Presentation(deck).slides)
    code = [i + 1 for i, s in enumerate(D) if any(p.text.startswith(">>> ") for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs)]; mans = [e["slide"] for e in ex]; print("manifest slides == code slides of", os.path.basename(deck), ":", mans == code, "(%d examples, %d code slides)" % (len(ex), len(code)))
    miss = []
    for e in ex:
        first = e["code"].splitlines()[0].strip(); tx = " ".join(p.text for sh in D[e["slide"] - 1].shapes if sh.has_text_frame for p in sh.text_frame.paragraphs)
        if first and first not in tx: miss.append((e["order"], e["slide"], first[:40]))
    print("examples whose first code line is not on their slide:", miss or "none"); return mans == code and not miss
a = sys.argv[1]; ok = {"plan": planc, "arcone": arcone, "codeslides": codeslides}[a](); print("RESULT %s: %s" % (a, "PASS" if ok else "FAIL")); sys.exit(0 if ok else 1)
