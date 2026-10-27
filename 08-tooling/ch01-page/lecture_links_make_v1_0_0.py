#!/usr/bin/env python3
"""lecture_links_make v1_0_0 (chapter 1, page 9.73.0): lecture_v3_1_1.json -> lecture_v3_1_2.json.
Over v3_1_1 only one thing changes: a slide of the deck whose Source hyperlinks sit beside case texts (slide 19) gets "links"
[{group, label, url, finds}], read from the deck v3_1_0 itself (group = the case title, label = the link text, finds = the case line above it,
url = the hyperlink address), for every deck hyperlink that the lecture does not yet carry. Nothing is typed by hand.
Usage: lecture_links_make_v1_0_0.py <deck.pptx> <lecture_in.json> <lecture_out.json>"""
import sys, json
from pptx import Presentation
deck, fin, fout = sys.argv[1:4]
lec = json.load(open(fin)); prs = Presentation(deck)
have = set()
for s in lec["slides"]:
    for l in s.get("links", []): have.add(l["url"])
    if s.get("read"): have.add(s["read"]["url"])
def runs(shapes):
    for x in shapes:
        if x.shape_type == 6: yield from runs(x.shapes)
        if x.has_text_frame:
            for p in x.text_frame.paragraphs:
                t = "".join(r.text for r in p.runs)
                links = [r.hyperlink.address for r in p.runs if r.hyperlink.address]
                if t.strip(): yield t.strip(), (links[0] if links else None)
added = {}
for i, sl in enumerate(prs.slides, 1):
    lines = list(runs(sl.shapes)); new = []
    for k, (t, u) in enumerate(lines):
        if u and u not in have and u not in {l["url"] for l in new}:
            new.append({"group": lines[k-2][0], "label": t, "url": u, "finds": lines[k-1][0]})
    if new:
        tgt = next(s for s in lec["slides"] if s["n"] == i); tgt.setdefault("links", []).extend(new); added[i] = len(new)
lec["_note"] += " | v3_1_2: slide link groups added from the deck %s by lecture_links_make_v1_0_0.py: %s" % (deck.rsplit("/", 1)[-1], json.dumps(added))
json.dump(lec, open(fout, "w"), indent=1, ensure_ascii=False)
print("lecture_links_make: %s -> %s; links added per slide %s" % (fin.rsplit("/", 1)[-1], fout.rsplit("/", 1)[-1], added))
