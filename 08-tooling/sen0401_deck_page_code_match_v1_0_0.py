#!/usr/bin/env python3
"""sen0401_deck_page_code_match v1_0_0 - the newest chapter deck and page: every code line of the chapter-code manifest is on the deck and in the page's strings (the owner's 'slide code found nowhere in the page' rule; the older measure counts lines
with '>>>' prompts, so on a deck without prompts it counts 0 lines and says nothing). Usage: NN  -> prints counts, exit 1 when a line is missing."""
import glob, json, re, sys
from pptx import Presentation
N = sys.argv[1]; R = "/home/claude/sen0401/03-materials/ch%s/" % N
k = lambda x: [int(v) for v in re.findall(r"_v(\d+)_(\d+)_(\d+)", x)[-1]]
deck = sorted(glob.glob(R + "SEN0401_Ch%s_*_3e_v*.pptx" % N), key=k)[-1]; page = sorted(glob.glob(R + "page/sen0401_ch%s_page_v*.html" % N), key=k)[-1]; man = sorted(glob.glob(R + "page/chapter_code_ch%s_v*.json" % N), key=k)[-1]
T = {ln.rstrip() for s in Presentation(deck).slides for sh in s.shapes if sh.has_text_frame for ln in sh.text_frame.text.split("\n")}
D = json.loads(re.search(r'<script[^>]*id="data"[^>]*>(.*?)</script>', open(page, encoding="utf8").read(), re.S).group(1))
def strs(o):
    if isinstance(o, str): yield o
    elif isinstance(o, dict):
        for v in o.values(): yield from strs(v)
    elif isinstance(o, list):
        for v in o: yield from strs(v)
P = {l.rstrip() for s in strs(D) for l in s.split("\n")}
lines = [(e.get("slide"), l.rstrip()) for e in json.load(open(man))["examples"] for l in e["code"].split("\n") if l.strip()]
nd = [x for x in lines if x[1] not in T]; npg = [x for x in lines if x[1] not in P]
print("deck", deck.split("/")[-1], "| page", page.split("/")[-1], "| manifest", man.split("/")[-1])
print("%d manifest code lines; not on the deck: %d; not in the page: %d" % (len(lines), len(nd), len(npg)))
print("RESULT:", "PASS" if not nd and not npg else "FAIL"); sys.exit(0 if not nd and not npg else 1)
