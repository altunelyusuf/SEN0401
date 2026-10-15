"""Chapter 14 post-render (the same fix as chapter 8) fix, run after `sen0401_chbuild render 14`: the shared renderer draws 'Questions?' on the closing slide as a box 8.8 in wide at y 4.82, which
overlaps the page number '26 / 26' (x 8.6, y 5.08) and is refused by the fit check (every earlier deck has the same overlap). This narrows that one box to 7.8 in. Nothing else changes.
Usage: post_render_v1_0_0.py <deck.pptx>"""
__version__ = "1.0.0"
import sys
from pptx import Presentation
from pptx.util import Inches
p = sys.argv[1]
prs = Presentation(p)
n = 0
for sh in prs.slides[len(prs.slides) - 1].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip() == "Questions?":
        sh.width = Inches(7.8); n += 1
assert n == 1, "the Questions? box was not found once"
prs.save(p)
print("closing slide: 'Questions?' box narrowed to 7.8 in")
