#!/usr/bin/env python3
"""Builds the five D9 fixtures of the deck standard check v1_2_0 (python-pptx). Usage: build_fixtures_v1_0_0.py OUTDIR"""
import os, sys
from pptx import Presentation
from pptx.dml.color import RGBColor
LIGHT, DARK = "FAF6EF", "14213D"
CASES = {  # name -> backgrounds of the 5 slides (L light, D dark); expected D9 in the file name
    "fx_light_only_PASS": "LLLLL",
    "fx_dark_last_only_PASS": "LLLLD",
    "fx_dark_middle_FAIL": "LLDLL",
    "fx_two_dark_FAIL": "LLLDD",
    "fx_dark_first_FAIL": "DLLLL",
}
out = sys.argv[1]; os.makedirs(out, exist_ok=True)
for name, pat in CASES.items():
    p = Presentation()
    from pptx.util import Inches
    for k, ch in enumerate(pat, 1):
        s = p.slides.add_slide(p.slide_layouts[6])
        s.shapes.add_textbox(Inches(8), Inches(7), Inches(1.5), Inches(0.4)).text_frame.text = "%d / %d" % (k, len(pat))
        f = s.background.fill; f.solid(); f.fore_color.rgb = RGBColor.from_string(DARK if ch == "D" else LIGHT)
    p.save(os.path.join(out, name + ".pptx"))
    print("built", name, pat)
