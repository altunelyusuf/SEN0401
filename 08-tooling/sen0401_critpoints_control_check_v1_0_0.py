#!/usr/bin/env python3
"""sen0401_critpoints_control_check v1_0_0 - the lineage agent's OWN control checks of the critical-points list (TK_Exam_CriticalPoints_Build). It shares no code with the executor's checker or reader.
Usage: sen0401_critpoints_control_check_v1_0_0.py KIND [COMMIT]   KIND: rerun | tamper | quotes | tokens | counts | newest | nochange COMMIT. Exit 0 only when met."""
import copy, glob, json, os, re, subprocess, sys, tempfile, zipfile
from pptx import Presentation
R = "/home/claude/sen0401/"; J = R + "08-tooling/exam/sen0401_exam_critical_points_v1_0_0.json"; CK = R + "08-tooling/exam/sen0401_exam_critical_points_check_v1_0_0.py"
d = json.load(open(J)); A = d["list_A_points_carried_by_a_source"]; B = d["list_B_not_carried_for_the_owner"]
N = lambda s: re.sub(r"\s+", " ", s.replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"').replace("–", "-").replace("—", "-").replace("\xa0", " ")).strip().lower()
def walk(shs):
    o = []
    for sh in shs:
        if sh.shape_type == 6: o += walk(sh.shapes)
        elif getattr(sh, "has_table", False) and sh.has_table: o += [c.text for r in sh.table.rows for c in r.cells]
        elif sh.has_text_frame: o.append(sh.text_frame.text)
    return o
def raw(path):
    z = zipfile.ZipFile(path); return N(" ".join(re.sub(r"<[^>]+>", " ", z.read(n).decode("utf8", "replace")) for n in z.namelist() if n.endswith(".xml")))
def fields(p):
    s = p["source"]; f = R + s["file"]
    if s["kind"] == "page": return {"page": N(open(f, encoding="utf8").read().replace('\\"', '"'))}
    if s["kind"] == "legacy2025": return {"slide": raw(f), "notes": raw(f)}
    sl = Presentation(f).slides[s["slide"] - 1]
    return {"slide": N(" ".join(walk(sl.shapes))), "notes": N(sl.notes_slide.notes_text_frame.text) if sl.has_notes_slide else ""}
def quotes():
    bad = []; per = {}
    for p in A:
        s = p["source"]; fl = fields(p); k = s["kind"]; per[k] = per.get(k, 0) + 1
        txt = fl.get(s["field"], fl.get("slide"))
        if not all(N(q) in txt for q in s["quote"].split("||")): bad.append(p["id"])
    print("quotes found by an independent reader: %d of %d points (%s); missing: %s" % (len(A) - len(bad), len(A), per, bad[:10])); return not bad
def tokens():
    bad = []; other = []; tr = re.compile(r"\d[\d,\.]*\d|\d|BIP\s?\d+|[A-Za-z]+_[A-Za-z_]+|0x[0-9a-f]+")
    for p in A:
        fl = fields(p); hay = " ".join(fl.values()) + " " + N(p["source"]["quote"])
        for t in tr.findall(p["point"]):
            t = t.strip(".,")
            if t and N(t) not in hay:
                f = R + p["source"]["file"]; whole = raw(f) if p["source"]["kind"] != "page" else hay
                (other if N(t) in whole else bad).append((p["id"], t))
    print("numbers and code-like tokens of the point statements: absent from the cited slide but present on another slide of the same file (a point may join two slides): %s; absent from the whole file: %d %s" % (other, len(bad), bad[:8])); return not bad
def rerun():
    ok = True
    for c in ("1-6", "1-7"):
        r = subprocess.run([sys.executable, "-B", CK, "--chapters", c], capture_output=True, text=True, cwd=R); line = [l for l in r.stdout.splitlines() if l.startswith("RESULT")]
        fails = [l.split()[1] for l in r.stdout.splitlines() if l.startswith("FAIL")]
        if fails == ["C8"]:   # C8 reads untracked files of the whole folder: the other executor (TK_Exam_PaperRules_Build) has an untracked file there; allowed only then, and printed
            ut = subprocess.run(["git", "-C", R, "ls-files", "--others", "--exclude-standard", "08-tooling/exam/"], capture_output=True, text=True).stdout.split()
            mine = [u for u in ut if "critical_points" in u or "pptx_text" in u]
            print("executor's check, chapters %s: only C8 differs; the untracked file(s) of ANOTHER task in 08-tooling/exam/: %s; untracked files of this task: %s; the clean 8 of 8 run is the one made before that file appeared" % (c, ut, mine)); ok &= not mine
        else:
            print("executor's check, chapters %s: exit %d, %s" % (c, r.returncode, line)); ok &= r.returncode == 0 and "PASS" in " ".join(line)
    return ok
def tamper():
    tmp = tempfile.mkdtemp(prefix="critctl_", dir="/tmp/claude-0/s"); res = []
    cases = {"quote": lambda t: t["list_A_points_carried_by_a_source"][3]["source"].__setitem__("quote", t["list_A_points_carried_by_a_source"][3]["source"]["quote"] + " zzz"),
             "number": lambda t: t["list_A_points_carried_by_a_source"][10].__setitem__("point", t["list_A_points_carried_by_a_source"][10]["point"] + " It takes 4242 rounds."),
             "slide": lambda t: t["list_A_points_carried_by_a_source"][20]["source"].__setitem__("slide", 9999),
             "listB": lambda t: t["list_B_not_carried_for_the_owner"][0].__setitem__("source", "x"),
             "chapter": lambda t: t.__setitem__("list_A_points_carried_by_a_source", [p for p in t["list_A_points_carried_by_a_source"] if p["chapter"] != 4])}
    for k, f in cases.items():
        t = copy.deepcopy(d); f(t); fp = os.path.join(tmp, k + ".json"); json.dump(t, open(fp, "w"))
        r = subprocess.run([sys.executable, "-B", CK, "--chapters", "1-7", "--json", fp], capture_output=True, text=True, cwd=R); res.append((k, r.returncode != 0)); print("tampered copy (%s): checker exit %d -> %s" % (k, r.returncode, "REFUSED" if r.returncode else "ACCEPTED"))
    import shutil; shutil.rmtree(tmp, ignore_errors=True); return all(x[1] for x in res)
def counts():
    from collections import Counter
    c = Counter(p["chapter"] for p in A); want = {1: 12, 2: 14, 3: 19, 4: 13, 5: 20, 6: 23, 7: 12}; ids = [p["id"] for p in A]
    ok = dict(c) == want and len(A) == 113 and len(B) == 8 and all("source" not in b and b["gap"] and b["asked_of_owner"] for b in B) and len(set(ids)) == len(ids) and d["default_chapters"] == [1, 2, 3, 4, 5, 6] and all(p["doubtful"] for p in A if p["chapter"] == 7) and not any(p["doubtful"] for p in A if p["chapter"] < 7)
    flags = [p["id"] for p in A if p["source"].get("slide") == 65 and p["chapter"] == 5 and p["basis"] == "EXAM"] + [p["id"] for p in A if p["basis"] == "OWN25"] + [p["id"] for p in A if p["basis"] == "PAGE"]
    print("per chapter %s (claimed %s); list B %d items without a source; explicit flags in the list: %s" % (dict(sorted(c.items())), want, len(B), flags)); return ok and len(flags) >= 3
def newest():
    bad = []
    for p in A:
        s = p["source"]
        if s["kind"] != "deck": continue
        fs = sorted(glob.glob(R + "03-materials/ch%02d/*.pptx" % p["chapter"]), key=lambda q: [int(x) for x in re.findall(r"_v(\d+)_(\d+)_(\d+)\.pptx$", q)[0]])
        if os.path.basename(fs[-1]) != os.path.basename(s["file"]): bad.append(p["id"])
    print("points whose deck is not the newest deck of its chapter: %s" % bad[:10]); return not bad
def nochange(c):
    st = subprocess.run(["git", "-C", R, "show", "--name-status", "--format=", c], capture_output=True, text=True).stdout.split("\n"); st = [l for l in st if l]
    notnew = [l for l in st if not l.startswith("A\t")]; out = [l for l in st if not l.split("\t")[1].startswith("08-tooling/exam/")]
    print("commit %s: %d files, all added: %s; all under 08-tooling/exam/: %s" % (c, len(st), not notnew, not out)); return st and not notnew and not out
k = sys.argv[1]
ok = {"rerun": rerun, "tamper": tamper, "quotes": quotes, "tokens": tokens, "counts": counts, "newest": newest, "nochange": lambda: nochange(sys.argv[2])}[k]()
print("RESULT", "MET" if ok else "NOT MET"); sys.exit(0 if ok else 1)
