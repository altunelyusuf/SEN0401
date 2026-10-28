#!/usr/bin/env python3
"""sen0401_deck06_control_check v1.0.0 - the control's OWN checks of the chapter 6 deck v2_2_0 (story ST_Deck06_Labels, criterion AC9_Deck06_Labels). A NEW tool; it edits nothing.
Usage: sen0401_deck06_control_check_v1_0_0.py KIND
  sha       the deck v2_1_0 has the same SHA-256 in the commit before the executor's work (parent of the first commit) as now, and no later commit touches it
  parts     same zip part list; the differing parts are exactly ppt/charts/chart6.xml (the chart of slide 51) and docProps/core.xml
  chartdiff chart6.xml: after replacing the number format code 0.## by General in the data labels the part equals the old one except the properties are not in it (byte for byte); values, chart type and showVal unchanged
  textdiff  the text of every shape and table and the notes of every slide are identical
  links     the set of hyperlink addresses of v2_1_0 equals that of v2_2_0 (count printed)
  render    the control's OWN rendering: soffice converts both decks to PDF, slide 51 is read with pdftotext -bbox; the count of intersecting word boxes among ALL words of the chart area (not only numbers, a wider method than the executor's) is 0 for v2_2_0, and the same method finds the crowding in v2_1_0 (> 0); the page images are kept in the verify folder
  dcheck    the deck check v1_1_0 under venv314 on v2_2_0: exit 0 and '0 mismatch(es)'
  arc0      the arc check v1_3_0 on v2_2_0 refuses nothing (exit 0)
Exit 0 only when the kind holds."""
__version__ = "1.0.0"
import hashlib, itertools, os, re, shutil, subprocess, sys, tempfile, zipfile
R = os.path.dirname(os.path.dirname(os.path.abspath(__file__))); D = os.path.join(R, "03-materials/ch06")
OLD = os.path.join(D, "SEN0401_Ch06_Transactions_3e_v2_1_0.pptx"); NEW = os.path.join(D, "SEN0401_Ch06_Transactions_3e_v2_2_0.pptx")
PARENT = "e8c31c7"; V = "/home/claude/venv314/bin/python"; KEEP = os.path.join(R, "08-tooling/verify/iter32/bridge_dk06")
def sha(b): return hashlib.sha256(b).hexdigest()
def git(*a): return subprocess.run(["git", "-C", R] + list(a), capture_output=True).stdout
def done(ok, msg): print(msg); print("RESULT " + ("PASS" if ok else "FAIL")); sys.exit(0 if ok else 1)
def texts(path):
    from pptx import Presentation
    out = []
    for s in Presentation(path).slides:
        l = []
        for sh in s.shapes:
            if sh.has_text_frame: l += [p.text for p in sh.text_frame.paragraphs]
            if sh.has_table: l += [" | ".join(c.text for c in r.cells) for r in sh.table.rows]
        out.append((l, s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""))
    return out
def links(p):
    z = zipfile.ZipFile(p); o = set()
    for n in z.namelist():
        if re.match(r"ppt/slides/_rels/slide\d+\.xml\.rels$", n):
            for m in re.finditer(r"<Relationship [^>]*>", z.read(n).decode()):
                if "/hyperlink" in m.group(0): o.add(re.search(r'Target="([^"]+)"', m.group(0)).group(1))
    return o
def k_sha():
    rel = "03-materials/ch06/SEN0401_Ch06_Transactions_3e_v2_1_0.pptx"; b = git("show", "%s:%s" % (PARENT, rel)); n = open(os.path.join(R, rel), "rb").read(); ok = len(b) > 0 and sha(b) == sha(n)
    print("%s before(%s) %s after %s %s" % (rel, PARENT, sha(b), sha(n), "equal" if ok else "DIFFERENT")); ch = git("log", "--format=%h", "%s..HEAD" % PARENT, "--", rel).split(); print("commits since touching it: %s" % ch); done(ok and not ch, "sha")
def k_parts():
    a, b = zipfile.ZipFile(OLD), zipfile.ZipFile(NEW); na, nb = sorted(a.namelist()), sorted(b.namelist()); diff = sorted(n for n in na if n in nb and a.read(n) != b.read(n))
    rel = a.read("ppt/slides/_rels/slide51.xml.rels").decode(); print("slide 51 chart part: %s" % re.findall(r"charts/(chart\d+\.xml)", rel))
    print("part lists equal: %s (%d parts); differing parts: %s" % (na == nb, len(na), diff)); done(na == nb and diff == ["docProps/core.xml", "ppt/charts/chart6.xml"] and "charts/chart6.xml" in rel, "parts")
def k_chartdiff():
    a = zipfile.ZipFile(OLD).read("ppt/charts/chart6.xml").decode(); b = zipfile.ZipFile(NEW).read("ppt/charts/chart6.xml").decode()
    back = b.replace('<c:numFmt formatCode="0.##" sourceLinked="0"/>', '<c:numFmt formatCode="General" sourceLinked="0"/>'); vals = lambda x: re.findall(r"<c:v>([^<]*)</c:v>", x); typ = lambda x: sorted(set(re.findall(r"<c:(\w+Chart)>", x)))
    print("after restoring the old format code the part equals the old part byte for byte: %s; values equal %s (%d); types %s / %s; dLbls blocks old %d new %d; showVal=1 old %d new %d; format codes in new: %s" % (back == a, vals(a) == vals(b), len(vals(b)), typ(a), typ(b), a.count("<c:dLbls>"), b.count("<c:dLbls>"), a.count('<c:showVal val="1"/>'), b.count('<c:showVal val="1"/>'), sorted(set(re.findall(r'formatCode="([^"]*)"', b)))))
    done(back == a and vals(a) == vals(b) and typ(a) == typ(b) and b.count('<c:showVal val="1"/>') == a.count('<c:showVal val="1"/>') > 0, "chartdiff")
def k_text():
    a, b = texts(OLD), texts(NEW); ch = [i + 1 for i in range(min(len(a), len(b))) if a[i] != b[i]]; print("slides %d and %d; slides whose text or notes differ: %s" % (len(a), len(b), ch)); done(len(a) == len(b) and not ch, "textdiff")
def k_links():
    a, b = links(OLD), links(NEW); print("distinct addresses v2_1_0: %d, v2_2_0: %d; only in old: %s; only in new: %s" % (len(a), len(b), sorted(a - b), sorted(b - a))); done(a == b and len(b) > 0, "links")
def count(pdf):
    x = subprocess.run(["pdftotext", "-bbox", "-f", "51", "-l", "51", pdf, "-"], capture_output=True, text=True).stdout
    W, H = [float(v) for v in re.search(r'<page width="([\d.]+)" height="([\d.]+)"', x).groups()]
    ws = [(float(a), float(b), float(c), float(d), t) for a, b, c, d, t in re.findall(r'<word xMin="([\d.]+)" yMin="([\d.]+)" xMax="([\d.]+)" yMax="([\d.]+)">([^<]*)</word>', x)]
    ch = [w for w in ws if w[0] > .5 * W and .25 * H < w[1] < .80 * H]
    pairs = [(p[4], q[4]) for p, q in itertools.combinations(ch, 2) if min(p[2], q[2]) - max(p[0], q[0]) > 0 and min(p[3], q[3]) - max(p[1], q[1]) > 0]
    return len(ch), pairs
def k_render():
    tmp = tempfile.mkdtemp(prefix="sen0401_d06ctl_", dir=os.environ.get("TMPDIR")); os.makedirs(KEEP, exist_ok=True); res = {}
    try:
        for tag, src in (("v2_1_0", OLD), ("v2_2_0", NEW)):
            shutil.copy(src, os.path.join(tmp, tag + ".pptx"))
            subprocess.run(["soffice", "--headless", "-env:UserInstallation=file://" + tmp + "/prof", "--convert-to", "pdf", "--outdir", tmp, os.path.join(tmp, tag + ".pptx")], capture_output=True, text=True, timeout=900)
            pdf = os.path.join(tmp, tag + ".pdf"); res[tag] = count(pdf)
            subprocess.run(["pdftoppm", "-f", "51", "-l", "51", "-r", "80", "-png", pdf, os.path.join(KEEP, "control_slide51_" + tag)], capture_output=True)
            print("%s: words in the chart area %d, intersecting pairs %d %s" % (tag, res[tag][0], len(res[tag][1]), res[tag][1]))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    done("v2_1_0" in res and "v2_2_0" in res and len(res["v2_1_0"][1]) > 0 and len(res["v2_2_0"][1]) == 0, "render")
def k_dcheck():
    r = subprocess.run([V, "-B", os.path.join(R, "08-tooling/sen0401_deck_check_v1_1_0.py"), "06", NEW], capture_output=True, text=True); out = (r.stdout + r.stderr).strip(); print(out[-700:]); print("exit %d" % r.returncode); done(r.returncode == 0 and "0 mismatch(es)" in out, "dcheck")
def k_arc():
    r = subprocess.run([sys.executable, "-B", os.path.join(R, "08-tooling/sen0401_deck_arc_check_v1_3_0.py"), "06"], capture_output=True, text=True, env=dict(os.environ, SEN0401_DECK=NEW)); out = (r.stdout + r.stderr).strip(); print(out[-500:]); print("exit %d" % r.returncode); done(r.returncode == 0 and "REFUSED" not in out, "arc0")
if __name__ == "__main__":
    {"sha": k_sha, "parts": k_parts, "chartdiff": k_chartdiff, "textdiff": k_text, "links": k_links, "render": k_render, "dcheck": k_dcheck, "arc0": k_arc}[sys.argv[1]]()
